"""
TA Analyzer - shared residual store for FFT / LPSVD analyses.

After Global Analysis, LDA, or single-trace Kinetic Fit completes, the
residual data needed by the vibrational-coherence (FFT / LPSVD) workflow
is auto-saved here as an Excel file.  The Coherence dialog can then list
the folder and reload any saved residual on demand.

Two file kinds are supported:

* 2D residual matrix  ``R[lambda, t]``
  sheet ``Residual2D``: top row = delays, left column = wavelengths,
  interior cells = residual values.

* 1D residual trace  ``r(t)``
  sheet ``Residual1D``: column A = delay, column B = residual.

Every file carries a ``Meta`` sheet with key/value rows describing the
analysis type, dataset tag, target wavelength (if any), and the
timestamp the file was written at.
"""
from __future__ import annotations

import os
import re
import datetime as _dt
from typing import Any

import numpy as np


# ---------------------------------------------------------------------
# Folder layout
# ---------------------------------------------------------------------
DEFAULT_DIRNAME = 'residuals'


def get_residual_dir(base_dir: str | None = None) -> str:
    """Return the residuals folder (creating it if needed)."""
    if base_dir is None:
        base_dir = os.path.dirname(os.path.abspath(__file__))
    path = os.path.join(base_dir, DEFAULT_DIRNAME)
    os.makedirs(path, exist_ok=True)
    return path


# ---------------------------------------------------------------------
# Filename helpers
# ---------------------------------------------------------------------
_BAD_CHARS = re.compile(r'[^A-Za-z0-9._-]+')


def sanitize(token: str) -> str:
    s = _BAD_CHARS.sub('_', (token or '').strip())
    return s.strip('._-') or 'data'


def make_filename(analysis: str, dataset: str,
                  wavelength: float | None = None,
                  timestamp: _dt.datetime | None = None) -> str:
    """Build a filename: ``<Analysis>_<dataset>[_<wl>nm]_<TS>.xlsx``."""
    ts = (timestamp or _dt.datetime.now()).strftime('%Y%m%d-%H%M%S')
    parts = [sanitize(analysis), sanitize(dataset)]
    if wavelength is not None and np.isfinite(wavelength):
        parts.append(f'{wavelength:.1f}nm')
    parts.append(ts)
    return '_'.join(parts) + '.xlsx'


# ---------------------------------------------------------------------
# Write
# ---------------------------------------------------------------------
def _meta_rows(analysis: str, dataset: str,
               wavelength: float | None,
               time_unit: str | None,
               extra: dict[str, Any] | None,
               kind: str,
               timestamp: _dt.datetime) -> list[tuple[str, str]]:
    rows = [
        ('analysis', analysis),
        ('dataset', dataset or ''),
        ('kind', kind),
        ('time_unit', time_unit or ''),
        ('timestamp', timestamp.isoformat(timespec='seconds')),
    ]
    if wavelength is not None and np.isfinite(wavelength):
        rows.append(('wavelength_nm', f'{float(wavelength):.6g}'))
    else:
        rows.append(('wavelength_nm', ''))
    if extra:
        for k, v in extra.items():
            rows.append((str(k), '' if v is None else str(v)))
    return rows


def _open_wb():
    try:
        from openpyxl import Workbook
    except ImportError as e:
        raise RuntimeError(
            'openpyxl is required to save residual files. '
            'Install with: pip install openpyxl') from e
    return Workbook()


def save_residual_2d(analysis: str, dataset: str,
                     wl: np.ndarray, t: np.ndarray, R: np.ndarray,
                     *, base_dir: str | None = None,
                     wavelength: float | None = None,
                     time_unit: str | None = None,
                     extra: dict[str, Any] | None = None,
                     filename: str | None = None,
                     full_path: str | None = None) -> str:
    """Save a 2D residual matrix; return the path.

    If ``full_path`` is given, write to that exact location (parent
    directory created if missing).  Otherwise write to the residuals
    folder under an auto-generated filename.
    """
    wl = np.asarray(wl, dtype=float).ravel()
    t = np.asarray(t, dtype=float).ravel()
    R = np.asarray(R, dtype=float)
    if R.shape != (wl.size, t.size):
        raise ValueError(
            f'Residual shape {R.shape} does not match '
            f'({wl.size} wl, {t.size} t).')

    ts = _dt.datetime.now()
    if full_path is not None:
        path = full_path
        parent = os.path.dirname(path)
        if parent:
            os.makedirs(parent, exist_ok=True)
    else:
        folder = get_residual_dir(base_dir)
        if filename is None:
            filename = make_filename(analysis, dataset, wavelength, ts)
        path = os.path.join(folder, filename)

    wb = _open_wb()
    ws_meta = wb.active
    ws_meta.title = 'Meta'
    ws_meta.append(['key', 'value'])
    for k, v in _meta_rows(analysis, dataset, wavelength, time_unit,
                           extra, '2D', ts):
        ws_meta.append([k, v])

    ws = wb.create_sheet('Residual2D')
    ws.append([None] + [float(x) for x in t])
    for i in range(wl.size):
        row = [float(wl[i])]
        ri = R[i, :]
        for j in range(t.size):
            v = ri[j]
            row.append(float(v) if np.isfinite(v) else None)
        ws.append(row)

    wb.save(path)
    return path


def save_residual_1d(analysis: str, dataset: str,
                     t: np.ndarray, r: np.ndarray,
                     *, base_dir: str | None = None,
                     wavelength: float | None = None,
                     time_unit: str | None = None,
                     extra: dict[str, Any] | None = None,
                     filename: str | None = None,
                     full_path: str | None = None) -> str:
    """Save a 1D residual trace; return the path.

    If ``full_path`` is given, write to that exact location (parent
    directory created if missing).  Otherwise write to the residuals
    folder under an auto-generated filename.
    """
    t = np.asarray(t, dtype=float).ravel()
    r = np.asarray(r, dtype=float).ravel()
    if t.size != r.size:
        raise ValueError(
            f'1D residual length mismatch: t has {t.size}, r has {r.size}.')

    ts = _dt.datetime.now()
    if full_path is not None:
        path = full_path
        parent = os.path.dirname(path)
        if parent:
            os.makedirs(parent, exist_ok=True)
    else:
        folder = get_residual_dir(base_dir)
        if filename is None:
            filename = make_filename(analysis, dataset, wavelength, ts)
        path = os.path.join(folder, filename)

    wb = _open_wb()
    ws_meta = wb.active
    ws_meta.title = 'Meta'
    ws_meta.append(['key', 'value'])
    for k, v in _meta_rows(analysis, dataset, wavelength, time_unit,
                           extra, '1D', ts):
        ws_meta.append([k, v])

    ws = wb.create_sheet('Residual1D')
    ws.append(['delay', 'residual'])
    for ti, ri in zip(t, r):
        ws.append([float(ti),
                   float(ri) if np.isfinite(ri) else None])

    wb.save(path)
    return path


# ---------------------------------------------------------------------
# Read
# ---------------------------------------------------------------------
def _load_wb(path: str):
    try:
        from openpyxl import load_workbook
    except ImportError as e:
        raise RuntimeError(
            'openpyxl is required to read residual files. '
            'Install with: pip install openpyxl') from e
    return load_workbook(path, data_only=True, read_only=True)


def _read_meta(wb) -> dict[str, str]:
    meta: dict[str, str] = {}
    if 'Meta' not in wb.sheetnames:
        return meta
    ws = wb['Meta']
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row:
            continue
        k = row[0]
        v = row[1] if len(row) > 1 else None
        if k is None:
            continue
        meta[str(k)] = '' if v is None else str(v)
    return meta


_KFIT_CSV_WL_RE = re.compile(
    r'kfit[_-]?([0-9]+(?:\.[0-9]+)?)\s*nm', re.IGNORECASE)


def _load_residual_csv(path: str) -> dict[str, Any]:
    """Load a 2-column CSV residual (delay, residual) written by kfit.

    Wavelength is best-effort extracted from the filename pattern
    ``kfit_{wl}nm_residual.csv`` and folded into ``meta`` for the
    coherence UI's status label.
    """
    # numpy handles both '# ...' comment prefixes (comments='#') and a
    # naked header line — try both.
    try:
        mat = np.loadtxt(path, delimiter=',', comments='#', skiprows=1)
    except Exception:
        mat = np.loadtxt(path, delimiter=',', comments='#')
    mat = np.atleast_2d(mat)
    if mat.shape[1] < 2:
        raise ValueError(
            f'{path}: expected at least 2 columns (delay, residual); '
            f'got shape {mat.shape}.')
    t = np.asarray(mat[:, 0], dtype=float)
    R = np.asarray(mat[:, 1], dtype=float)
    meta: dict[str, str] = {'analysis': 'Kfit', 'kind': '1D'}
    m = _KFIT_CSV_WL_RE.search(os.path.basename(path))
    if m:
        meta['wavelength_nm'] = m.group(1)
    return {'kind': '1D', 'meta': meta, 't': t, 'wl': None, 'R': R}


def load_residual(path: str) -> dict[str, Any]:
    """Load a residual file written by save_residual_* or by Kfit CSV.

    Dispatches by extension:
      * ``.csv/.tsv/.txt/.dat`` — 2-column ``(delay, residual)`` layout
        (1D only; produced by the single-trace Kinetic Fit dialog)
      * ``.xlsx`` — the Meta+Residual{1D,2D} workbook layout produced by
        Global Analysis / LDA (2D) and the legacy Kfit auto-save (1D)

    Returns a dict with keys:
        kind     '2D' or '1D'
        meta     dict[str, str]
        t        ndarray of delays
        wl       ndarray of wavelengths (2D only)
        R        ndarray, shape (wl, t) for 2D or (t,) for 1D

    The underlying workbook is closed before returning so the file is
    not held open on Windows.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext in ('.csv', '.tsv', '.txt', '.dat'):
        return _load_residual_csv(path)
    wb = _load_wb(path)
    try:
        meta = _read_meta(wb)
        kind = meta.get('kind') or (
            '2D' if 'Residual2D' in wb.sheetnames else
            '1D' if 'Residual1D' in wb.sheetnames else '')

        if kind == '2D':
            if 'Residual2D' not in wb.sheetnames:
                raise ValueError(f'{path}: missing Residual2D sheet.')
            ws = wb['Residual2D']
            rows = list(ws.iter_rows(values_only=True))
            if not rows:
                raise ValueError(f'{path}: Residual2D sheet is empty.')
            header = rows[0]
            t = np.array([float(x) for x in header[1:] if x is not None],
                         dtype=float)
            wl_list, R_rows = [], []
            for row in rows[1:]:
                if row is None or row[0] is None:
                    continue
                wl_list.append(float(row[0]))
                vals = []
                for j in range(t.size):
                    idx = 1 + j
                    v = row[idx] if idx < len(row) else None
                    vals.append(float(v) if v is not None else np.nan)
                R_rows.append(vals)
            wl = np.asarray(wl_list, dtype=float)
            R = np.asarray(R_rows, dtype=float)
            return {'kind': '2D', 'meta': meta, 't': t, 'wl': wl, 'R': R}

        if kind == '1D':
            if 'Residual1D' not in wb.sheetnames:
                raise ValueError(f'{path}: missing Residual1D sheet.')
            ws = wb['Residual1D']
            ts, rs = [], []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i == 0:
                    continue
                if row is None or row[0] is None:
                    continue
                ts.append(float(row[0]))
                v = row[1] if len(row) > 1 else None
                rs.append(float(v) if v is not None else np.nan)
            t = np.asarray(ts, dtype=float)
            R = np.asarray(rs, dtype=float)
            return {'kind': '1D', 'meta': meta, 't': t, 'wl': None, 'R': R}

        raise ValueError(
            f'{path}: unknown residual kind (meta.kind={kind!r}).')
    finally:
        try:
            wb.close()
        except Exception:
            pass


# ---------------------------------------------------------------------
# Folder listing
# ---------------------------------------------------------------------
def list_residual_files(base_dir: str | None = None) -> list[str]:
    """Return a list of xlsx file paths in the residuals folder.

    Newer files (by mtime) first.
    """
    folder = get_residual_dir(base_dir)
    out = []
    for name in os.listdir(folder):
        if not name.lower().endswith('.xlsx'):
            continue
        out.append(os.path.join(folder, name))
    out.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    return out
